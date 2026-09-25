import io

from docx import Document

from app.services.docx_gen import quote_docx
from tests.fakes import make_client
from tests.sample_data import COMPANY, QUOTE


def _doc():
    return Document(io.BytesIO(quote_docx(QUOTE, COMPANY)))


def _all_text(doc):
    paras = [p.text for p in doc.paragraphs]
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join(paras + cells)


def test_docx_contains_header_and_footer():
    text = _all_text(_doc())
    for s in ["Quotation", "ETS/SS10/26-27", "Yale Electronics Services Pvt Ltd",
              "Mr Jayashekar R Yale", "23-05-2026", "Terms and Conditions:",
              "No.190/3,Kalkere Village", "GST NO: 29AADCE7362F1ZI", "Authorised Signatory"]:
        assert s in text, s


def test_docx_items_table_money_and_blank_rows():
    doc = _doc()
    items = next(t for t in doc.tables if t.rows[0].cells[0].text == "Sl.No.")
    assert items.rows[1].cells[3].text == "75,200.00"
    assert items.rows[2].cells[2].text == ""  # label-only row
    assert items.rows[-1].cells[4].text == "21,78,000.00"
    assert "None" not in _all_text(doc)


def test_docx_merged_group_cells():
    doc = _doc()
    items = next(t for t in doc.tables if t.rows[0].cells[0].text == "Sl.No.")
    # rows 3 and 4 (sl 3, 4) share one merged total cell
    assert items.rows[3].cells[4]._tc is items.rows[4].cells[4]._tc
    assert items.rows[3].cells[4].text == "21,02,800.00"


def test_docx_endpoint():
    c, _ = make_client()
    cid = c.post("/api/v1/customers", json={"company_name": "Yale"}).json()["id"]
    qid = c.post("/api/v1/quotes", json={"ref_no": "ETS/SS10/26-27", "customer_id": cid,
                                         "quote_date": "2026-05-23"}).json()["id"]
    r = c.get(f"/api/v1/quotes/{qid}/docx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert 'Quote-ETS-SS10-26-27-Yale.docx' in r.headers["content-disposition"]

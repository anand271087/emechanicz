"""Letterhead layout: logo header and company footer on every page; ref/issue/date flush right."""
import io
import re

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Pt

from app.services.docx_gen import quote_docx
from app.services.pdf import render_quote_html
from tests.sample_data import COMPANY, QUOTE

LONG_QUOTE = {**QUOTE, "quote_items": [
    {"sl_no": n, "description": f"Item {n} with a reasonably long description line", "qty": 1,
     "unit_price": 1000.0, "total": 1000.0, "group_id": None} for n in range(1, 71)]}


def _pdf_pages(quote):
    from pypdf import PdfReader
    from app.services.pdf import quote_pdf
    return PdfReader(io.BytesIO(quote_pdf(quote, COMPANY))).pages


@pytest.mark.pdf
def test_pdf_long_quote_spans_several_pages():
    assert len(_pdf_pages(LONG_QUOTE)) >= 2


@pytest.mark.pdf
def test_pdf_footer_on_every_page():
    for page in _pdf_pages(LONG_QUOTE):
        text = page.extract_text()
        assert "GST NO: 29AADCE7362F1ZI" in text
        assert "Emechanicz Test Solutions Pvt Ltd," in text


@pytest.mark.pdf
def test_pdf_logo_drawn_on_every_page():
    # Images are stored once per document, so count the draw operator on each page instead.
    for page in _pdf_pages(LONG_QUOTE):
        assert len(re.findall(rb"/\S+ Do\b", page.get_contents().get_data())) == 1


@pytest.mark.pdf
def test_pdf_footer_sits_at_page_bottom_even_on_short_quote():
    from pypdf import PdfReader
    from app.services.pdf import quote_pdf
    page = PdfReader(io.BytesIO(quote_pdf(QUOTE, COMPANY))).pages[0]
    height = float(page.mediabox.height)
    ys = []

    def visit(text, cm, tm, _font, _size):
        if "GST NO" in text:
            ys.append(cm[3] * tm[5] + cm[5])  # page y, measured up from the bottom edge

    page.extract_text(visitor_text=visit)
    assert ys and 0 < ys[0] < height * 0.1


def _info_rows(html):
    table = re.search(r'<table class="info">(.*?)</table>', html, re.S).group(1)
    return re.findall(r"<tr>(.*?)</tr>", table, re.S)


def test_html_ref_no_on_customer_line_then_issue_then_date():
    rows = _info_rows(render_quote_html(QUOTE, COMPANY))
    assert "Customer :" in rows[0] and "ETS Ref No : ETS/SS10/26-27" in rows[0]
    assert "Kind Attn :" in rows[1] and "Issue Status : 1.1" in rows[1]
    assert "Date : 23-05-2026" in rows[2]


def test_html_right_column_is_flush_right():
    css = re.search(r"\.info \.right \{([^}]*)\}", render_quote_html(QUOTE, COMPANY)).group(1)
    assert "text-align: right" in css


def _docx():
    return Document(io.BytesIO(quote_docx(QUOTE, COMPANY)))


def test_docx_logo_in_page_header_right_aligned():
    header = _docx().sections[0].header
    para = header.paragraphs[0]
    assert para.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert para._p.xpath(".//pic:pic")


def test_docx_company_footer_on_every_page():
    section = _docx().sections[0]
    assert not section.different_first_page_header_footer
    text = "\n".join(p.text for p in section.footer.paragraphs)
    assert "Emechanicz Test Solutions Pvt Ltd," in text and "GST NO: 29AADCE7362F1ZI" in text
    body = "\n".join(p.text for p in _docx().paragraphs)
    assert "GST NO:" not in body


def test_docx_ref_issue_date_flush_right():
    # One paragraph per line, right text after a right-aligned tab stop at the right margin:
    # works in Word, Google Docs and Pages without relying on table column widths.
    paras = _docx().paragraphs
    start = next(i for i, p in enumerate(paras) if p.text.startswith("Customer :"))
    lines = paras[start:start + 3]
    assert [p.text for p in lines] == [
        "Customer : Yale Electronics Services Pvt Ltd\tETS Ref No : ETS/SS10/26-27",
        "Kind Attn : Mr Jayashekar R Yale\tIssue Status : 1.1",
        "Dear Sir,\tDate : 23-05-2026",
    ]
    for p in lines:
        (stop,) = p.paragraph_format.tab_stops
        assert stop.alignment == WD_TAB_ALIGNMENT.RIGHT and stop.position == Pt(503)

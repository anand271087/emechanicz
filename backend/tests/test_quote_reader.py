"""Reading existing quotations (PDF, Excel, Word) into editable quote drafts."""
import io
import shutil
import subprocess

import pytest
from openpyxl import Workbook

from app.services.docx_gen import quote_docx
from app.services.quote_reader import parse_number, read_quote
from scripts.render_sample import COMPANY, sample_quote


def yale_xlsx(merge_groups=True) -> bytes:
    """An Excel sheet laid out like the original Yale quotation."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Quotation"])
    ws.append(["Customer : Yale Electronics Services Pvt Ltd"])
    ws.append(["Kind Attn : Mr Jayashekar R Yale", None, None, "ETS Ref No :", "ETS/SS10/26-27"])
    ws.append([None, None, None, "Issue Status :", "1.1"])
    ws.append(["Dear Sir,", None, None, "Date :", "23-05-2026"])
    ws.append(["Further to your enquiry, please find the quote below"])
    ws.append(["Sl.No.", "Description", "Qty", "Unit Price", "TotalPrice"])
    rows = [(1, "Rack with MS Heavy Duty Extruded Structure", 1, 75200, 75200),
            (2, "Sotware Develpoment /LabView Programming", 1, 450000, 450000),
            (3, "Labview installer professional", None, None, None),
            (4, "Documentation & Training", "", "", "")]
    for r in rows:
        ws.append(list(r))
    first = ws.max_row - 2  # row of item 2
    if merge_groups:
        for col in "CDE":
            ws.merge_cells(f"{col}{first}:{col}{first + 1}")
    ws.append([None, None, None, None, 525200])
    ws.append(["INR", "Five Lakh Twenty Five Thousand Two Hundred Only"])
    ws.append(["Terms and Conditions:"])
    ws.append([1, "Delivery: 10-12 Weeks from the date of PO & Confirmation"])
    ws.append([2, "Order to be placed on:"])
    ws.append([None, "Emechanicz Test Solutions Pvt Ltd ,"])
    ws.append([3, "Validity of quotation: 15 Days from the date of proposal."])
    ws.append(["Thanking you and assuring you of our best services at all the times."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def maxeye_xlsx(total=1654200) -> bytes:
    """MaxEye-style labels: no spaces, double colons, Issue Ref, S.NO / Price (INR) headers."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Customer :MaxEye Technologies Pvt. Ltd..", None, None, "Ref  No: ETS/P301/26-27"])
    ws.append(["Kind Attn :: Mr.Reegan.M", None, None, "Issue Ref: 1.1"])
    ws.append([None, None, None, "Date:01/09/2026"])
    ws.append(["S.NO", "Description", "Qty", "Price (INR)", "Total (INR)"])
    ws.append([1, "Pnematic Fixture (1IN1 PCBA )", 12, "1,37,850.00", "16,54,200.00"])
    ws.append([None, None, None, None, total])
    ws.append(["INR", "Sixteen Lakh Fifty Four Thoasand Two Hundred Only"])
    ws.append([1, "Delivery:3-4weeks from the date of PO & Confirmation"])
    ws.append([2, "Payment Terms: Net 30 Days"])
    ws.append(["For EMechanicZ Test Solution Pvt ltd ,"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- numbers ---

@pytest.mark.parametrize("text,value", [
    ("1,37,850.00", 137850.0), ("7 5,200.0", 75200.0), ("1 6,54,200.00", 1654200.0),
    ("₹ 40,000", 40000.0), ("12", 12.0), (75200, 75200.0), ("", None), (None, None), ("abc", None)])
def test_parse_number(text, value):
    assert parse_number(text) == value


# --- Excel ---

def test_xlsx_header_fields():
    d = read_quote("Quote.xlsx", yale_xlsx())
    assert (d["customer_name"], d["kind_attn"], d["ref_no"], d["issue_status"], d["quote_date"]) == (
        "Yale Electronics Services Pvt Ltd", "Mr Jayashekar R Yale", "ETS/SS10/26-27", "1.1", "2026-05-23")
    assert d["intro"] == "Dear Sir,\nFurther to your enquiry, please find the quote below"


def test_xlsx_items_with_label_row_and_merged_price():
    items = read_quote("Quote.xlsx", yale_xlsx())["items"]
    assert [(i["sl_no"], i["qty"], i["unit_price"], i["join_above"]) for i in items] == [
        (1, 1, 75200, False), (2, 1, 450000, False), (3, None, None, True), (4, None, None, False)]
    assert items[0]["description"] == "Rack with MS Heavy Duty Extruded Structure"


def test_xlsx_total_and_terms():
    d = read_quote("Quote.xlsx", yale_xlsx())
    assert d["printed_total"] == 525200 and d["warnings"] == []
    assert d["terms"] == ["Delivery: 10-12 Weeks from the date of PO & Confirmation",
                          "Order to be placed on:\nEmechanicz Test Solutions Pvt Ltd ,",
                          "Validity of quotation: 15 Days from the date of proposal."]


def test_maxeye_style_labels():
    d = read_quote("Quote.xlsx", maxeye_xlsx())
    assert (d["customer_name"], d["kind_attn"], d["ref_no"], d["issue_status"], d["quote_date"]) == (
        "MaxEye Technologies Pvt. Ltd.", "Mr.Reegan.M", "ETS/P301/26-27", "1.1", "2026-09-01")
    assert [(i["qty"], i["unit_price"]) for i in d["items"]] == [(12, 137850)]
    assert d["terms"] == ["Delivery:3-4weeks from the date of PO & Confirmation", "Payment Terms: Net 30 Days"]


def test_total_mismatch_is_warned():
    d = read_quote("Quote.xlsx", maxeye_xlsx(total=1700000))
    assert any("don't add up" in w for w in d["warnings"])


def test_xls_old_excel_format():
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Quote")
    cells = [["Customer : Tata Elxsi", "", "", "Ref No : ETS/P9/26-27"], ["Kind Attn : Mr A", "", "", "Date : 15/08/2026"],
             ["Sl.No.", "Description", "Qty", "Unit Price", "Total"], [1, "Relay card", 2, 66000, 132000],
             ["", "", "", "", 132000]]
    for r, row in enumerate(cells):
        for c, v in enumerate(row):
            ws.write(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    d = read_quote("old.xls", buf.getvalue())
    assert (d["customer_name"], d["ref_no"], d["quote_date"]) == ("Tata Elxsi", "ETS/P9/26-27", "2026-08-15")
    assert [(i["description"], i["qty"], i["unit_price"]) for i in d["items"]] == [("Relay card", 2, 66000)]


# --- PDF and Word (the app's own exports use the same layout as the originals) ---

@pytest.mark.pdf
def test_pdf_round_trip_of_full_yale_quote():
    from app.services.pdf import quote_pdf
    d = read_quote("Quote.pdf", quote_pdf(sample_quote(), COMPANY))
    assert (d["customer_name"], d["kind_attn"], d["ref_no"], d["quote_date"]) == (
        "Yale Electronics Services Pvt Ltd", "Mr Jayashekar R Yale", "ETS/SS10/26-27", "2026-05-23")
    assert len(d["items"]) == 20
    assert [i["sl_no"] for i in d["items"] if i["join_above"]] == [12, 19]
    assert d["items"][1]["description"].startswith("Rack (Monitor Arm,PDU")
    assert d["printed_total"] == 2138000 and d["warnings"] == []
    assert len(d["terms"]) == 7 and d["terms"][4].startswith("Order to be placed on:\nEmechanicz")


def test_docx_round_trip():
    d = read_quote("Quote.docx", quote_docx(sample_quote(), COMPANY))
    assert (d["customer_name"], d["ref_no"], d["quote_date"], d["issue_status"]) == (
        "Yale Electronics Services Pvt Ltd", "ETS/SS10/26-27", "2026-05-23", "1.1")
    assert len(d["items"]) == 20 and d["printed_total"] == 2138000
    assert [i["sl_no"] for i in d["items"] if i["join_above"]] == [12, 19]
    assert d["warnings"] == []


@pytest.mark.skipif(not (shutil.which("antiword") and shutil.which("textutil")),
                    reason="needs antiword (server) and textutil (macOS) to build a .doc")
def test_doc_old_word_format(tmp_path):
    src = tmp_path / "q.docx"
    src.write_bytes(quote_docx(sample_quote(), COMPANY))
    subprocess.run(["textutil", "-convert", "doc", str(src), "-output", str(tmp_path / "q.doc")], check=True)
    d = read_quote("q.doc", (tmp_path / "q.doc").read_bytes())
    assert d["ref_no"] == "ETS/SS10/26-27" and len(d["items"]) == 20


# --- problems ---

def test_unsupported_type_rejected():
    with pytest.raises(ValueError, match="PDF, Excel"):
        read_quote("notes.txt", b"hello")


def test_corrupt_file_rejected():
    with pytest.raises(ValueError, match="couldn't be opened"):
        read_quote("broken.xlsx", b"PK\x03\x04 nope")


def test_missing_fields_are_warned():
    wb = Workbook()
    wb.active.append(["Sl.No.", "Description", "Qty", "Unit Price", "Total"])
    wb.active.append([1, "Fan", 1, 100, 100])
    buf = io.BytesIO()
    wb.save(buf)
    d = read_quote("bare.xlsx", buf.getvalue())
    assert {"No customer name found", "No ref no found", "No date found"} <= set(d["warnings"])


def test_no_item_table_warned():
    wb = Workbook()
    wb.active.append(["Customer : Yale"])
    buf = io.BytesIO()
    wb.save(buf)
    assert "No item table found" in read_quote("x.xlsx", buf.getvalue())["warnings"]


def test_wrapped_description_lines_join_the_row_above_for_old_word_files():
    from app.services.quote_reader import _parse
    table = [["Sl.No.", "Description", "Qty", "Unit Price", "TotalPrice"],
             ["1", "Rack with MS Heavy Duty", "1", "75,200.00", "75,200.00"],
             ["", "Extruded Structure 25U x", "", "", ""],
             ["", "600W x 1000mmD", "", "", ""],
             ["2", "Documentation & Training", "", "", ""],
             ["", "", "", "", "75,200.00"]]
    items = _parse(["Customer : Yale | Ref No : R1"], [table], continuations=True)["items"]
    assert [(i["sl_no"], i["description"]) for i in items] == [
        (1, "Rack with MS Heavy Duty Extruded Structure 25U x 600W x 1000mmD"), (2, "Documentation & Training")]


def test_rows_without_numbers_stay_separate_for_other_formats():
    from app.services.quote_reader import _parse
    table = [["Sl.No.", "Description", "Qty", "Unit Price", "Total"], ["1", "Rack", "1", "5", "5"], ["", "Loose note", "", "", ""]]
    assert len(_parse([], [table])["items"]) == 2

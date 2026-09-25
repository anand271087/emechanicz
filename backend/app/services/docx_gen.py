import io

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.services.formatting import ORDINAL_RE, format_date, inr, qty
from app.services.pdf import STATIC_DIR
from app.services.quote_layout import table_rows

NAVY = RGBColor(0x17, 0x36, 0x5D)
HEADER_FILL = "A6C9EC"
COL_WIDTHS = [Pt(38), Pt(291), Pt(27), Pt(76), Pt(71)]
LEFT, CENTER, RIGHT = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT


def _shade(cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _write(cell, text: str, align=LEFT, bold=False) -> None:
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _para(doc, text="", size=None, bold=False, align=LEFT, color=None, indent=None, space_before=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    if indent is not None:
        p.paragraph_format.left_indent = indent
    if text:
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
    return p


def _set_widths(table, widths) -> None:
    """Fixed layout with explicit grid so every renderer honours the widths."""
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(int(w.pt * 20) for w in widths)))
    tbl_w.set(qn("w:type"), "dxa")
    for grid_col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        grid_col.set(qn("w:w"), str(int(w.pt * 20)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def _borders(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    table._tbl.tblPr.append(borders)


def _header(doc, quote: dict) -> None:
    t = doc.add_table(rows=1, cols=2)
    _set_widths(t, [Pt(366), Pt(137)])
    title = t.cell(0, 0).paragraphs[0].add_run("Quotation")
    title.bold, title.font.size = True, Pt(18)
    t.cell(0, 0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    logo = t.cell(0, 1).paragraphs[0]
    logo.alignment = CENTER
    logo.add_run().add_picture(str(STATIC_DIR / "logo.png"), width=Pt(52))

    customer = (quote.get("customers") or {}).get("company_name", "")
    intro = (quote.get("intro") or "").split("\n")
    info = doc.add_table(rows=4, cols=3)
    _set_widths(info, [Pt(293), Pt(62), Pt(148)])
    lines = [
        (f"Customer : {customer}", "", ""),
        (f"Kind Attn : {quote.get('kind_attn', '')}", "ETS Ref No :", quote["ref_no"]),
        ("", "Issue Status :", quote.get("issue_status", "")),
        (intro[0], "Date :", format_date(quote["quote_date"])),
    ]
    for r, (left, label, value) in enumerate(lines):
        info.cell(r, 0).paragraphs[0].add_run(left)
        info.cell(r, 1).paragraphs[0].add_run(label)
        info.cell(r, 1).paragraphs[0].alignment = RIGHT
        info.cell(r, 2).paragraphs[0].add_run(value)
    for line in intro[1:]:
        _para(doc, line, size=10.8)


def _items(doc, quote: dict) -> None:
    rows = table_rows(quote.get("quote_items", []))
    t = doc.add_table(rows=len(rows) + 2, cols=5)
    t.style = "Table Grid"
    _borders(t)
    _set_widths(t, COL_WIDTHS)
    for i, h in enumerate(["Sl.No.", "Description", "Qty", "Unit Price", "TotalPrice"]):
        cell = t.cell(0, i)
        _write(cell, h, CENTER, bold=True)
        _shade(cell, HEADER_FILL)
    for n, r in enumerate(rows, start=1):
        _write(t.cell(n, 0), str(r["sl_no"]), CENTER)
        _write(t.cell(n, 1), r["description"])
        if r["span"] == 0:
            continue
        for col, text, align in ((2, qty(r["qty"]), CENTER), (3, inr(r["unit_price"]), RIGHT),
                                 (4, inr(r["total"]), RIGHT)):
            cell = t.cell(n, col)
            if r["span"] > 1:
                cell = cell.merge(t.cell(n + r["span"] - 1, col))
            _write(cell, text, align)
    _write(t.cell(len(rows) + 1, 4), inr(quote.get("subtotal")), RIGHT, bold=True)

    words = doc.add_table(rows=1, cols=2)
    _set_widths(words, [Pt(39), Pt(464)])
    _write(words.cell(0, 0), "INR", CENTER)
    _write(words.cell(0, 1), quote.get("amount_in_words", ""))


def _terms(doc, quote: dict, company: dict) -> None:
    terms = quote.get("terms") or []
    if not terms:
        return
    _para(doc, "Terms and Conditions:", bold=True, space_before=24)
    t = doc.add_table(rows=0, cols=2)
    for i, term in enumerate(terms, start=1):
        first, *extra = term.split("\n")
        row = t.add_row()
        _write(row.cells[0], str(i), RIGHT)
        _write(row.cells[1], first)
        for line in extra:
            row = t.add_row()
            _write(row.cells[1], line, bold=True)
    _set_widths(t, [Pt(35), Pt(468)])
    if company.get("closing_line"):
        _para(doc, company["closing_line"], indent=Pt(39))


def _footer(doc, company: dict) -> None:
    _para(doc, f"For {company.get('signatory_name') or company.get('name', '')} ,", space_before=22)
    _para(doc)
    _para(doc, "Authorised Signatory")
    if company.get("system_generated_note"):
        _para(doc, company["system_generated_note"], indent=Pt(57))
    _para(doc, f"{company.get('name', '')},", size=9.3, align=CENTER, color=NAVY, space_before=14)

    addr = _para(doc, align=CENTER)
    pos, text = 0, company.get("footer_address", "")
    for m in ORDINAL_RE.finditer(text):
        for chunk, sup in ((text[pos:m.start(2)], False), (m.group(2), True)):
            run = addr.add_run(chunk)
            run.font.size, run.font.color.rgb, run.font.superscript = Pt(9), NAVY, sup
        pos = m.end(2)
    run = addr.add_run(text[pos:])
    run.font.size, run.font.color.rgb = Pt(9), NAVY

    _para(doc, f"M {company.get('phone', '')} - E: {company.get('emails', '')}", size=9.9,
          align=CENTER, color=NAVY)
    _para(doc, f"GST NO: {company.get('gst_no', '')}", size=9, align=CENTER,
          color=RGBColor(0x0E, 0x28, 0x41))


def quote_docx(quote: dict, company: dict) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Pt(46)
    section.top_margin, section.bottom_margin = Pt(34), Pt(36)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(9.9)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.1

    _header(doc, quote)
    _para(doc)
    _items(doc, quote)
    _terms(doc, quote, company)
    _footer(doc, company)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

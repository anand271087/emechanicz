"""Read existing quotations (PDF, Excel, Word) into editable quote drafts.

Every format is reduced to the same two things:
  lines  – the document's text, one string per line, cells joined with " | "
  tables – grids of cell text, where None marks a cell covered by a merged cell above
and one parser pulls the quote out of those.
"""
import io
import re
import shutil
import subprocess
import tempfile
from datetime import date, datetime
from pathlib import Path

Cell = str | None
Grid = list[list[Cell]]

SUPPORTED = (".pdf", ".xlsx", ".xls", ".docx", ".doc")
UNSUPPORTED = "Upload a PDF, Excel (.xlsx, .xls) or Word (.docx, .doc) file."


# ---------- small helpers ----------

def parse_number(value) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = re.sub(r"(?i)₹|rs\.?|inr|,|\s", "", str(value))
    try:
        return float(text) if text else None
    except ValueError:
        return None


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%Y")
    if isinstance(value, date):
        return value.strftime("%d-%m-%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).split())


def _join(cells) -> str:
    return " | ".join(c for c in cells if c)


def _parse_date(text: str) -> str | None:
    m = re.search(r"(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})", text or "")
    if not m:
        return None
    d, mth, y = (int(g) for g in m.groups())
    y = y + 2000 if y < 100 else y
    try:
        return date(y, mth, d).isoformat()
    except ValueError:
        return None


# ---------- extractors: bytes -> (lines, tables) ----------

def _from_pdf(content: bytes) -> tuple[list[str], list[Grid]]:
    import pdfplumber

    pages_lines, tables = [], []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            pages_lines.append([" ".join(l.split()) for l in (page.extract_text() or "").splitlines() if l.strip()])
            for t in page.extract_tables():
                tables.append([[None if c is None else " ".join(c.split()) for c in row] for row in t])
    if len(pages_lines) > 1:  # drop running headers/footers repeated on every page
        repeated = set.intersection(*(set(p) for p in pages_lines))
        pages_lines = [[l for l in p if l not in repeated] for p in pages_lines]
    return [l for p in pages_lines for l in p], tables


def _grid_from_sheet(rows: list[list], covered: set[tuple[int, int]]) -> Grid:
    return [[None if (r, c) in covered else _cell_text(v) for c, v in enumerate(row)] for r, row in enumerate(rows)]


def _from_xlsx(content: bytes) -> tuple[list[str], list[Grid]]:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(content), data_only=True).worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    covered = {(r - 1, c - 1) for rng in ws.merged_cells.ranges
               for r in range(rng.min_row, rng.max_row + 1) for c in range(rng.min_col, rng.max_col + 1)
               if (r, c) != (rng.min_row, rng.min_col)}
    grid = _grid_from_sheet(rows, covered)
    return [_join(r) for r in grid], [grid]


def _from_xls(content: bytes) -> tuple[list[str], list[Grid]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content, formatting_info=True)
    sh = book.sheet_by_index(0)
    rows = []
    for r in range(sh.nrows):
        row = []
        for c in range(sh.ncols):
            cell = sh.cell(r, c)
            row.append(xlrd.xldate_as_datetime(cell.value, book.datemode) if cell.ctype == xlrd.XL_CELL_DATE else cell.value)
        rows.append(row)
    covered = {(r, c) for rlo, rhi, clo, chi in sh.merged_cells
               for r in range(rlo, rhi) for c in range(clo, chi) if (r, c) != (rlo, clo)}
    grid = _grid_from_sheet(rows, covered)
    return [_join(r) for r in grid], [grid]


def _from_docx(content: bytes) -> tuple[list[str], list[Grid]]:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(content))
    lines, tables = [], []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            for text in Paragraph(child, doc).text.split("\n"):
                if text.strip():
                    lines.append(" | ".join(" ".join(p.split()) for p in text.split("\t") if p.strip()))
        elif child.tag == qn("w:tbl"):
            grid, above = [], {}
            for row in Table(child, doc).rows:
                cells, seen = [], set()
                for c, cell in enumerate(row.cells):
                    if id(cell._tc) in seen:  # repeated by a horizontal merge
                        continue
                    seen.add(id(cell._tc))
                    cells.append(None if above.get(c) is cell._tc else " ".join(cell.text.split()))
                    above[c] = cell._tc
                grid.append(cells)
                lines.append(_join(grid[-1]))
            tables.append(grid)
    return lines, tables


def _from_doc(content: bytes) -> tuple[list[str], list[Grid]]:
    if not shutil.which("antiword"):
        raise ValueError("Old Word (.doc) files can't be read here. Save it as .docx or PDF and upload again.")
    with tempfile.NamedTemporaryFile(suffix=".doc") as f:
        f.write(content)
        f.flush()
        out = subprocess.run(["antiword", "-w", "0", f.name], capture_output=True, text=True, timeout=60)
    if out.returncode != 0:
        raise ValueError("This Word file couldn't be opened. Check it opens in Word, then try again.")
    lines, tables, current = [], [], None
    for raw in out.stdout.splitlines():
        if raw.strip().startswith("|"):
            cells = [" ".join(c.split()) for c in raw.strip().strip("|").split("|")]
            if current is None:
                current = []
                tables.append(current)
            current.append(cells)
            lines.append(_join(cells))
        else:
            current = None
            if raw.strip():
                lines.append(" | ".join(" ".join(p.split()) for p in re.split(r"\t| {3,}", raw) if p.strip()))
    return lines, tables


READERS = {".pdf": _from_pdf, ".xlsx": _from_xlsx, ".xls": _from_xls, ".docx": _from_docx, ".doc": _from_doc}


# ---------- parser: (lines, tables) -> draft ----------

_END = r"(?=\s*\||\s+(?:ETS\s+)?Ref\.?\s*No|\s+Kind\s+Attn|\s+Issue\s|\s+Date\s*:|$)"
FIELDS = {
    "customer_name": re.compile(r"Customer\s*:+\s*(?:\|\s*)?(?P<v>.+?)" + _END, re.I),
    "kind_attn": re.compile(r"Kind\s*Attn\.?\s*:+\s*(?:\|\s*)?(?P<v>.+?)" + _END, re.I),
    "ref_no": re.compile(r"Ref\.?\s*No\.?\s*:+\s*(?:\|\s*)?(?P<v>[A-Za-z0-9][A-Za-z0-9/._-]*)", re.I),
    "issue_status": re.compile(r"Issue\s*(?:Status|Ref|No)?\.?\s*:+\s*(?:\|\s*)?(?P<v>\d+(?:\.\d+)?)", re.I),
    "date": re.compile(r"\bDate\s*:+\s*(?:\|\s*)?(?P<v>\d{1,2}[./-]\d{1,2}[./-]\d{2,4})", re.I),
}
TERM = re.compile(r"^(\d{1,2})\s*[.)]?\s*(?:\|\s*)?(?P<t>\S.*)$")
STOP = re.compile(r"^(Thanking|For\s+E\s*Mechanic|Authori[sz]ed|This is System)", re.I)
HEADER_WORDS = {"sl": r"^(s\.?\s*no|sl|sr)", "desc": r"desc", "qty": r"^(qty|quantity|nos?)\b",
                "price": r"(unit|price|rate)", "total": r"(total|amount)"}


def _find_columns(row: list[Cell]) -> dict[str, int] | None:
    texts = [(c or "").lower() for c in row]
    cols: dict[str, int] = {}
    for key in ("desc", "qty", "total", "price", "sl"):
        for i, t in enumerate(texts):
            if i in cols.values() or not t:
                continue
            if key == "price" and "total" in t:
                continue
            if re.search(HEADER_WORDS[key], t):
                cols[key] = i
                break
    return cols if {"desc", "qty"} <= cols.keys() else None


def _items(tables: list[Grid], continuations: bool = False) -> tuple[list[dict], float | None, bool]:
    """Line items, printed grand total, and whether an item table was found.

    continuations: a row with no Sl.No. and no numbers continues the description above
    (old Word files, where long cell text is wrapped onto several table lines).
    """
    for grid in tables:
        for h, header in enumerate(grid):
            cols = _find_columns(header)
            if cols:
                break
        else:
            continue
        get = lambda row, key: row[cols[key]] if key in cols and cols[key] < len(row) else ""
        items, total = [], None
        for row in grid[h + 1:]:
            desc, sl = get(row, "desc"), get(row, "sl")
            qty_c, price_c, total_c = get(row, "qty"), get(row, "price"), get(row, "total")
            first = next((c for c in row if c), "")
            if first.upper().startswith("INR") or re.match(r"terms", first, re.I):
                break
            if not desc:
                if desc is None and items:  # description merged downwards: nothing new
                    continue
                if parse_number(total_c) is not None and not sl:
                    total = parse_number(total_c)
                    break
                continue
            qty, price, line_total = parse_number(qty_c), parse_number(price_c), parse_number(total_c)
            if continuations and items and not sl and qty is None and price is None and line_total is None:
                items[-1]["description"] += " " + desc
                continue
            join_above = bool(items) and qty_c is None and price_c is None and total_c is None
            if qty is None and price is None and line_total is not None:
                qty, price = 1.0, line_total
            elif qty is not None and price is None and line_total is not None:
                price = line_total / qty if qty else None
            elif (qty is None) != (price is None):
                qty = price = None
            sl_no = parse_number(sl)
            items.append({"sl_no": int(sl_no) if sl_no else len(items) + 1, "description": desc,
                          "qty": qty, "unit_price": price, "join_above": join_above})
        return items, total, True
    return [], None, False


def _table_start(lines: list[str]) -> int:
    for i, line in enumerate(lines):
        if re.search(r"desc", line, re.I) and re.search(r"\bqty\b|quantity", line, re.I):
            return i
    return len(lines)


def _intro(header_lines: list[str]) -> str:
    for i, line in enumerate(header_lines):
        if re.match(r"^Dear\b", line, re.I):
            first = re.split(r"\s*\|\s*|\s+Date\s*:", line, maxsplit=1)[0].strip()
            out = [first]
            nxt = header_lines[i + 1] if i + 1 < len(header_lines) else ""
            if nxt and not any(p.search(nxt) for p in FIELDS.values()) and not _parse_date(nxt):
                out.append(nxt.split(" | ")[0])
            return "\n".join(out)
    return ""


def _terms(lines: list[str]) -> list[str]:
    start = next((i for i, l in enumerate(lines) if re.match(r"^INR\b", l, re.I)), None)
    if start is None:
        start = next((i for i, l in enumerate(lines) if re.match(r"^Terms\b", l, re.I)), None)
    if start is None:
        return []
    terms: list[str] = []
    for line in lines[start + 1:]:
        if STOP.match(line):
            break
        if re.match(r"^Terms\b", line, re.I):
            continue
        m = TERM.match(line)
        text = line.replace(" | ", " ").strip()
        if m:
            terms.append(re.sub(r"^\d+(?=[A-Za-z])", "", m.group("t").replace(" | ", " ")).strip())
        elif terms:
            terms[-1] += "\n" + text
    return terms


def _parse(lines: list[str], tables: list[Grid], continuations: bool = False) -> dict:
    head = lines[:_table_start(lines)]
    found: dict[str, str] = {}
    for key, pattern in FIELDS.items():
        for line in head or lines:
            m = pattern.search(line)
            if m:
                found[key] = m.group("v").strip(" .,|") if key != "customer_name" else m.group("v").strip(" ,|")
                break
    customer = re.sub(r"\.{2,}", ".", found.get("customer_name", "")).strip()
    quote_date = _parse_date(found.get("date", "")) or next((d for d in map(_parse_date, head) if d), None)

    items, printed_total, has_table = _items(tables, continuations)
    computed = round(sum(i["qty"] * i["unit_price"] for i in items if i["qty"] is not None and i["unit_price"] is not None), 2)

    warnings = []
    if not customer:
        warnings.append("No customer name found")
    if not found.get("ref_no"):
        warnings.append("No ref no found")
    if not quote_date:
        warnings.append("No date found")
    if not has_table:
        warnings.append("No item table found")
    elif not items:
        warnings.append("The item table is empty")
    if printed_total is not None and abs(printed_total - computed) > 1:
        warnings.append(f"Line items don't add up to the printed total (items ₹{computed:,.2f}, "
                        f"printed ₹{printed_total:,.2f}). Check the rows.")

    return {"customer_name": customer, "kind_attn": found.get("kind_attn", ""), "ref_no": found.get("ref_no", ""),
            "issue_status": found.get("issue_status", "1.1"), "quote_date": quote_date, "intro": _intro(head),
            "items": items, "terms": _terms(lines), "printed_total": printed_total,
            "computed_total": computed, "warnings": warnings}


def read_quote(filename: str, content: bytes) -> dict:
    """Draft quote read from an uploaded file. Raises ValueError with a message for people."""
    ext = Path(filename or "").suffix.lower()
    if ext not in READERS:
        raise ValueError(UNSUPPORTED)
    try:
        lines, tables = READERS[ext](content)
    except ValueError:
        raise
    except Exception:
        raise ValueError("This file couldn't be opened. Check it opens normally on your computer, then try again.")
    return _parse(lines, tables, continuations=ext == ".doc")

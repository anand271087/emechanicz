"""Read item descriptions from an uploaded Excel (.xlsx) sheet."""
import io
import zipfile

from fastapi import HTTPException
from openpyxl import load_workbook

MAX_ROWS = 10_000
NOT_EXCEL = "Upload an Excel file (.xlsx). You can start from the template on the Items page."


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read_descriptions(content: bytes) -> tuple[list[str], int]:
    """Descriptions from the first sheet, and how many blank rows were skipped.

    Uses the column headed "Description" when the first non-empty row has one,
    otherwise the first column of every row.
    """
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError):
        raise HTTPException(400, NOT_EXCEL)
    rows = wb.worksheets[0].iter_rows(values_only=True)

    column, texts, blanks, started = 0, [], 0, False
    for i, row in enumerate(rows):
        if i >= MAX_ROWS:
            raise HTTPException(400, f"The sheet has more than {MAX_ROWS:,} rows. Split it into smaller files.")
        cells = [_cell_text(v) for v in row]
        if not started:
            if not any(cells):
                continue  # blank rows above the data
            started = True
            headers = [c.lower() for c in cells]
            if "description" in headers:
                column = headers.index("description")
                continue
        text = cells[column] if column < len(cells) else ""
        if text:
            texts.append(text)
        else:
            blanks += 1
    wb.close()
    if not texts:
        raise HTTPException(400, "No descriptions found. Put them under a \"Description\" heading in the first sheet.")
    return texts, blanks
